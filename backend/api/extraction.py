"""Text extraction for uploaded documents."""
import logging
import os

import docx
from pypdf import PdfReader

logger = logging.getLogger(__name__)

MIME_TYPES = {
    '.pdf': 'application/pdf',
    '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    '.txt': 'text/plain',
}


def get_extension(filename):
    return os.path.splitext(filename or '')[1].lower()


def extract_text(file, extension):
    """
    Extract text from an open binary file.

    Returns a (text, pages) tuple. `pages` is None when it does not apply.
    Raises ValueError when the file cannot be read as the given type.
    """
    file.seek(0)
    try:
        if extension == '.pdf':
            reader = PdfReader(file)
            text = '\n'.join((page.extract_text() or '') for page in reader.pages)
            pages = len(reader.pages)
        elif extension == '.docx':
            document = docx.Document(file)
            text = '\n'.join(paragraph.text for paragraph in document.paragraphs)
            # Estimate pages (rough calculation)
            pages = max(1, len(document.paragraphs) // 20)
        elif extension == '.txt':
            text = file.read().decode('utf-8', errors='replace')
            pages = None
        else:
            raise ValueError(f'Unsupported file type: {extension}')
    except ValueError:
        raise
    except Exception as exc:
        logger.warning('Failed to extract text from %s file: %s', extension, exc)
        raise ValueError('The file could not be read. It may be corrupted or password protected.') from exc
    finally:
        file.seek(0)

    # PostgreSQL text columns cannot store NUL characters
    return text.replace('\x00', '').strip(), pages
