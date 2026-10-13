import io
import tempfile
from unittest import mock

import docx
from django.core.cache import cache
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from rest_framework import status
from rest_framework.test import APITestCase
from django.test import SimpleTestCase

from .models import Document, Flashcard, StudySession, Summary, User
from .services import AIServiceError, AIServiceNotConfigured, OpenAIService

PASSWORD = 'correct-horse-battery'

SUMMARY_RESULT = {'full_summary': 'A summary.', 'key_points': ['One', 'Two']}
FLASHCARD_RESULT = [
    {'question': 'Q1?', 'answer': 'A1', 'category': 'Biology'},
    {'question': 'Q2?', 'answer': 'A2', 'category': 'General'},
]


def make_user(email, **extra):
    return User.objects.create_user(
        email=email, username=email.split('@')[0], password=PASSWORD, **extra
    )


def text_file(name='notes.txt', content=b'Cells are the basic unit of life.'):
    return SimpleUploadedFile(name, content, content_type='text/plain')


def minimal_pdf(text):
    """Build a tiny one-page PDF containing `text`."""
    stream = f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode()
    objects = [
        b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] '
        b'/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>',
        b'<< /Length %d >>\nstream\n%s\nendstream' % (len(stream), stream),
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
    ]
    pdf = b'%PDF-1.4\n'
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b'%d 0 obj\n%s\nendobj\n' % (number, body)
    xref_offset = len(pdf)
    pdf += b'xref\n0 %d\n0000000000 65535 f \n' % (len(objects) + 1)
    for offset in offsets:
        pdf += b'%010d 00000 n \n' % offset
    pdf += b'trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n' % (
        len(objects) + 1, xref_offset,
    )
    return pdf


class BaseAPITestCase(APITestCase):
    def setUp(self):
        cache.clear()  # throttle counters live in the cache
        media_root = tempfile.TemporaryDirectory()
        self.addCleanup(media_root.cleanup)
        media_override = override_settings(MEDIA_ROOT=media_root.name)
        media_override.enable()
        self.addCleanup(media_override.disable)

        self.user = make_user('alice@example.com')
        self.other = make_user('bob@example.com')
        self.client.force_authenticate(self.user)

    def make_document(self, user=None, text='Some text'):
        user = user or self.user
        return Document.objects.create(
            user=user, title='Doc', file=text_file(), file_size=10,
            file_type='text/plain', extracted_text=text,
        )


class AuthTests(BaseAPITestCase):
    def setUp(self):
        super().setUp()
        self.client.force_authenticate(None)

    def register(self, email, **extra):
        data = {
            'email': email, 'first_name': 'Test', 'last_name': 'User',
            'password': PASSWORD, 'password_confirm': PASSWORD, **extra,
        }
        return self.client.post('/api/auth/register/', data)

    def test_register_and_login(self):
        self.assertEqual(self.register('carol@example.com').status_code, status.HTTP_201_CREATED)

        response = self.client.post(
            '/api/auth/login/', {'email': 'carol@example.com', 'password': PASSWORD}
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('access', response.data)
        self.assertIn('refresh', response.data)
        self.assertEqual(response.data['user']['email'], 'carol@example.com')

    def test_register_generates_unique_username(self):
        # alice@example.com already owns the username "alice"
        response = self.register('alice@other.org')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['user']['username'], 'alice2')

    def test_register_rejects_weak_password(self):
        response = self.register(
            'dave@example.com', password='12345678', password_confirm='12345678'
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('password', response.data)

    def test_login_with_wrong_password_fails(self):
        response = self.client.post(
            '/api/auth/login/', {'email': 'alice@example.com', 'password': 'nope'}
        )
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_refresh_rotates_token(self):
        login = self.client.post(
            '/api/auth/login/', {'email': 'alice@example.com', 'password': PASSWORD}
        )
        response = self.client.post('/api/auth/refresh/', {'refresh': login.data['refresh']})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertIn('access', response.data)
        self.assertNotEqual(response.data['refresh'], login.data['refresh'])

        # The old refresh token is blacklisted after rotation
        reused = self.client.post('/api/auth/refresh/', {'refresh': login.data['refresh']})
        self.assertEqual(reused.status_code, status.HTTP_401_UNAUTHORIZED)

    def test_endpoints_require_authentication(self):
        for url in ('/api/documents/', '/api/flashcards/', '/api/dashboard/stats/'):
            self.assertEqual(self.client.get(url).status_code, status.HTTP_401_UNAUTHORIZED)

    def test_login_is_rate_limited(self):
        with mock.patch('api.throttles.AuthRateThrottle.get_rate', return_value='3/min'):
            codes = [
                self.client.post(
                    '/api/auth/login/', {'email': 'alice@example.com', 'password': 'nope'}
                ).status_code
                for _ in range(4)
            ]
        self.assertEqual(codes[-1], status.HTTP_429_TOO_MANY_REQUESTS)


class ProfileTests(BaseAPITestCase):
    def test_profile_update(self):
        response = self.client.patch('/api/auth/profile/', {'first_name': 'Alice'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.user.refresh_from_db()
        self.assertEqual(self.user.first_name, 'Alice')

    def test_user_cannot_make_themselves_admin(self):
        response = self.client.patch(
            '/api/auth/profile/', {'is_staff': True, 'is_superuser': True}, format='json'
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.user.refresh_from_db()
        self.assertFalse(self.user.is_staff)
        self.assertFalse(self.user.is_superuser)


class DocumentTests(BaseAPITestCase):
    def test_upload_extracts_text(self):
        response = self.client.post(
            '/api/documents/', {'title': 'Notes', 'file': text_file()}, format='multipart'
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertNotIn('extracted_text', response.data)
        self.assertNotIn('file', response.data)

        document = Document.objects.get(pk=response.data['id'])
        self.assertEqual(document.user, self.user)
        self.assertEqual(document.file_type, 'text/plain')
        self.assertEqual(document.extracted_text, 'Cells are the basic unit of life.')

    def test_upload_extracts_text_from_docx(self):
        buffer = io.BytesIO()
        word_document = docx.Document()
        word_document.add_paragraph('Mitochondria produce energy.')
        word_document.add_paragraph('Ribosomes build proteins.')
        word_document.save(buffer)
        upload = SimpleUploadedFile('cells.docx', buffer.getvalue())

        response = self.client.post(
            '/api/documents/', {'title': 'Cells', 'file': upload}, format='multipart'
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['pages'], 1)
        document = Document.objects.get(pk=response.data['id'])
        self.assertEqual(
            document.extracted_text, 'Mitochondria produce energy.\nRibosomes build proteins.'
        )
        self.assertIn('wordprocessingml', document.file_type)

    def test_upload_extracts_text_from_pdf(self):
        upload = SimpleUploadedFile('notes.pdf', minimal_pdf('Photosynthesis makes glucose'))

        response = self.client.post(
            '/api/documents/', {'title': 'Plants', 'file': upload}, format='multipart'
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['pages'], 1)
        self.assertEqual(response.data['file_type'], 'application/pdf')
        document = Document.objects.get(pk=response.data['id'])
        self.assertIn('Photosynthesis makes glucose', document.extracted_text)

    def test_upload_rejects_unsupported_type(self):
        for name in ('malware.exe', 'legacy.doc'):
            response = self.client.post(
                '/api/documents/', {'title': 'Bad', 'file': text_file(name)}, format='multipart'
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertIn('file', response.data)

    @override_settings(MAX_UPLOAD_SIZE=8)
    def test_upload_rejects_oversized_file(self):
        response = self.client.post(
            '/api/documents/', {'title': 'Big', 'file': text_file()}, format='multipart'
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn('file', response.data)

    def test_upload_rejects_corrupt_pdf(self):
        response = self.client.post(
            '/api/documents/',
            {'title': 'Broken', 'file': text_file('broken.pdf', b'not a pdf')},
            format='multipart',
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_upload_without_file_is_a_validation_error(self):
        response = self.client.post('/api/documents/', {'title': 'Empty'}, format='multipart')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_documents_are_private(self):
        document = self.make_document(self.other)
        self.assertEqual(self.client.get('/api/documents/').data['count'], 0)
        self.assertEqual(
            self.client.get(f'/api/documents/{document.pk}/').status_code,
            status.HTTP_404_NOT_FOUND,
        )
        self.assertEqual(
            self.client.delete(f'/api/documents/{document.pk}/').status_code,
            status.HTTP_404_NOT_FOUND,
        )

    def test_delete_returns_204_and_removes_file(self):
        document = self.make_document()
        storage, name = document.file.storage, document.file.name
        self.assertTrue(storage.exists(name))

        response = self.client.delete(f'/api/documents/{document.pk}/')
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(storage.exists(name))

    def test_delete_all_only_touches_own_documents(self):
        self.make_document()
        self.make_document(self.other)
        response = self.client.delete('/api/documents/delete_all/')
        self.assertEqual(response.data['deleted_count'], 1)
        self.assertEqual(Document.objects.filter(user=self.other).count(), 1)

    def test_page_size_can_be_raised_to_fetch_everything(self):
        for _ in range(25):
            self.make_document()
        self.assertEqual(len(self.client.get('/api/documents/').data['results']), 20)
        self.assertEqual(
            len(self.client.get('/api/documents/?page_size=100').data['results']), 25
        )


@mock.patch('api.views.OpenAIService')
class AIGenerationTests(BaseAPITestCase):
    def test_generate_summary_then_regenerate_updates_in_place(self, service):
        service.return_value.generate_summary.return_value = SUMMARY_RESULT
        document = self.make_document()
        url = f'/api/documents/{document.pk}/generate_summary/'

        first = self.client.post(url)
        self.assertEqual(first.status_code, status.HTTP_201_CREATED)
        self.assertEqual(first.data['key_points'], ['One', 'Two'])

        service.return_value.generate_summary.return_value = {
            'full_summary': 'A better summary.', 'key_points': [],
        }
        second = self.client.post(url)
        self.assertEqual(second.status_code, status.HTTP_200_OK)
        self.assertEqual(second.data['id'], first.data['id'])
        self.assertEqual(Summary.objects.filter(document=document).count(), 1)
        self.assertEqual(Summary.objects.get().full_summary, 'A better summary.')

    def test_generate_summary_for_other_users_document_is_404(self, service):
        document = self.make_document(self.other)
        response = self.client.post(f'/api/documents/{document.pk}/generate_summary/')
        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        service.assert_not_called()

    def test_generate_summary_without_text_is_400(self, service):
        document = self.make_document(text='')
        document.file.delete(save=True)
        response = self.client.post(f'/api/documents/{document.pk}/generate_summary/')
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_older_documents_fall_back_to_reading_the_file(self, service):
        service.return_value.generate_summary.return_value = SUMMARY_RESULT
        document = self.make_document(text='')

        response = self.client.post(f'/api/documents/{document.pk}/generate_summary/')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        document.refresh_from_db()
        self.assertEqual(document.extracted_text, 'Cells are the basic unit of life.')

    def test_ai_failure_does_not_leak_details(self, service):
        service.return_value.generate_summary.side_effect = AIServiceError('secret internals')
        document = self.make_document()

        response = self.client.post(f'/api/documents/{document.pk}/generate_summary/')
        self.assertEqual(response.status_code, status.HTTP_502_BAD_GATEWAY)
        self.assertNotIn('secret internals', str(response.data))
        self.assertNotIn('detail', response.data)

    def test_missing_api_key_is_503(self, service):
        service.side_effect = AIServiceNotConfigured('no key')
        document = self.make_document()
        response = self.client.post(f'/api/documents/{document.pk}/generate_flashcards/')
        self.assertEqual(response.status_code, status.HTTP_503_SERVICE_UNAVAILABLE)

    def test_generate_flashcards_from_document(self, service):
        service.return_value.generate_flashcards.return_value = FLASHCARD_RESULT
        document = self.make_document()

        response = self.client.post(
            f'/api/documents/{document.pk}/generate_flashcards/', {'num_cards': 2}
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(len(response.data), 2)
        self.assertTrue(all(card['id'] for card in response.data))
        self.assertEqual(Flashcard.objects.filter(user=self.user, document=document).count(), 2)

    def test_generate_flashcards_from_summary(self, service):
        service.return_value.generate_flashcards.return_value = FLASHCARD_RESULT
        document = self.make_document()
        summary = Summary.objects.create(user=self.user, document=document, full_summary='S')

        response = self.client.post(f'/api/summaries/{summary.pk}/generate_flashcards/')
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(Flashcard.objects.filter(summary=summary).count(), 2)

    def test_ai_endpoints_are_rate_limited(self, service):
        service.return_value.generate_summary.return_value = SUMMARY_RESULT
        document = self.make_document()
        url = f'/api/documents/{document.pk}/generate_summary/'

        with mock.patch('api.throttles.AIRateThrottle.get_rate', return_value='2/hour'):
            codes = [self.client.post(url).status_code for _ in range(3)]
        self.assertEqual(codes[-1], status.HTTP_429_TOO_MANY_REQUESTS)


class FlashcardTests(BaseAPITestCase):
    def make_card(self, user=None, **extra):
        return Flashcard.objects.create(
            user=user or self.user, question='Q?', answer='A', **extra
        )

    def test_create_returns_full_card_with_id(self):
        response = self.client.post(
            '/api/flashcards/', {'question': 'Q?', 'answer': 'A', 'category': 'Math'}
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertIsNotNone(response.data['id'])
        self.assertEqual(response.data['mastery_level'], 0)
        self.assertEqual(Flashcard.objects.get(pk=response.data['id']).user, self.user)

    def test_cannot_link_card_to_another_users_document_or_summary(self):
        document = self.make_document(self.other)
        summary = Summary.objects.create(user=self.other, document=document, full_summary='S')

        for field, target in (('document', document), ('summary', summary)):
            response = self.client.post(
                '/api/flashcards/', {'question': 'Q?', 'answer': 'A', field: target.pk}
            )
            self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
            self.assertIn(field, response.data)
        self.assertEqual(Flashcard.objects.count(), 0)

    def test_can_link_card_to_own_document(self):
        document = self.make_document()
        response = self.client.post(
            '/api/flashcards/', {'question': 'Q?', 'answer': 'A', 'document': document.pk}
        )
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data['document']['id'], document.pk)

    def test_update_cannot_change_progress_or_summary(self):
        other_document = self.make_document(self.other)
        other_summary = Summary.objects.create(
            user=self.other, document=other_document, full_summary='S'
        )
        card = self.make_card()

        response = self.client.put(
            f'/api/flashcards/{card.pk}/',
            {
                'question': 'New?', 'answer': 'New', 'category': 'General',
                'mastery_level': 100, 'review_count': 99, 'summary': other_summary.pk,
            },
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        card.refresh_from_db()
        self.assertEqual(card.question, 'New?')
        self.assertEqual(card.mastery_level, 0)
        self.assertEqual(card.review_count, 0)
        self.assertIsNone(card.summary)

    def test_delete_returns_204(self):
        card = self.make_card()
        response = self.client.delete(f'/api/flashcards/{card.pk}/')
        self.assertEqual(response.status_code, status.HTTP_204_NO_CONTENT)
        self.assertEqual(response.content, b'')

    def test_cards_are_private(self):
        card = self.make_card(self.other)
        self.assertEqual(self.client.get('/api/flashcards/').data['count'], 0)
        for method in (self.client.get, self.client.delete):
            self.assertEqual(
                method(f'/api/flashcards/{card.pk}/').status_code, status.HTTP_404_NOT_FOUND
            )

    def test_review_moves_mastery_in_steps(self):
        card = self.make_card()
        url = f'/api/flashcards/{card.pk}/review/'

        response = self.client.post(url, {'result': 'mastered'})
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data['mastery_level'], 25)
        self.assertEqual(response.data['review_count'], 1)
        self.assertIsNotNone(response.data['last_reviewed'])

        self.assertEqual(self.client.post(url, {'result': 'needs_practice'}).data['mastery_level'], 10)
        self.assertEqual(self.client.post(url, {'result': 'needs_practice'}).data['mastery_level'], 0)

        for _ in range(5):
            response = self.client.post(url, {'result': 'mastered'})
        self.assertEqual(response.data['mastery_level'], 100)

    def test_review_requires_valid_result(self):
        card = self.make_card()
        response = self.client.post(f'/api/flashcards/{card.pk}/review/', {'mastery_level': 100})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)

    def test_filters(self):
        document = self.make_document()
        self.make_card(category='Math', document=document)
        self.make_card(category='Biology')

        self.assertEqual(self.client.get('/api/flashcards/?category=Math').data['count'], 1)
        self.assertEqual(
            self.client.get(f'/api/flashcards/?document={document.pk}').data['count'], 1
        )


class SummaryTests(BaseAPITestCase):
    def test_filter_by_document_and_privacy(self):
        mine = self.make_document()
        Summary.objects.create(user=self.user, document=mine, full_summary='Mine')
        Summary.objects.create(user=self.user, document=self.make_document(), full_summary='Other doc')
        theirs = self.make_document(self.other)
        Summary.objects.create(user=self.other, document=theirs, full_summary='Theirs')

        self.assertEqual(self.client.get('/api/summaries/').data['count'], 2)
        response = self.client.get(f'/api/summaries/?document={mine.pk}')
        self.assertEqual(response.data['count'], 1)
        self.assertEqual(response.data['results'][0]['full_summary'], 'Mine')
        self.assertEqual(self.client.get(f'/api/summaries/?document={theirs.pk}').data['count'], 0)


class StudySessionAndDashboardTests(BaseAPITestCase):
    def test_study_sessions_feed_dashboard_study_time(self):
        for seconds in (3600, 1500):
            response = self.client.post(
                '/api/study-sessions/',
                {'duration_seconds': seconds, 'cards_reviewed': 10, 'cards_mastered': 4},
            )
            self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        StudySession.objects.create(user=self.other, duration_seconds=9999)
        Flashcard.objects.create(user=self.user, question='Q', answer='A', mastery_level=50)

        stats = self.client.get('/api/dashboard/stats/').data
        self.assertEqual(stats['study_seconds'], 5100)
        self.assertEqual(stats['study_time'], '1h 25m')
        self.assertEqual(stats['flashcards_count'], 1)
        self.assertEqual(stats['mastery'], 50)

    def test_study_session_rejects_bad_duration(self):
        response = self.client.post('/api/study-sessions/', {'duration_seconds': -5})
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)


class AdminTests(BaseAPITestCase):
    def test_admin_endpoints_reject_regular_users(self):
        for url in (
            '/api/admin/dashboard/stats/', '/api/admin/users/',
            f'/api/admin/users/{self.other.pk}/', '/api/admin/content/stats/',
        ):
            self.assertEqual(self.client.get(url).status_code, status.HTTP_403_FORBIDDEN)

    def test_staff_can_read_but_not_grant_access(self):
        staff = make_user('staff@example.com', is_staff=True)
        self.client.force_authenticate(staff)

        self.assertEqual(self.client.get('/api/admin/users/').data['count'], 3)
        self.assertEqual(self.client.get('/api/admin/content/stats/').status_code, status.HTTP_200_OK)

        response = self.client.patch(
            f'/api/admin/users/{self.other.pk}/', {'is_staff': True}, format='json'
        )
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.other.refresh_from_db()
        self.assertFalse(self.other.is_staff)

    def test_superuser_can_grant_staff_but_not_superuser(self):
        root = make_user('root@example.com', is_staff=True, is_superuser=True)
        self.client.force_authenticate(root)

        response = self.client.patch(
            f'/api/admin/users/{self.other.pk}/',
            {'is_staff': True, 'is_superuser': True},
            format='json',
        )
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.other.refresh_from_db()
        self.assertTrue(self.other.is_staff)
        self.assertFalse(self.other.is_superuser)


@override_settings(OPENAI_API_KEY='test-key', AI_MAX_INPUT_CHARS=50)
class OpenAIServiceTests(SimpleTestCase):
    def make_service(self, content):
        service = OpenAIService()
        service.client = mock.Mock()
        message = mock.Mock(content=content)
        service.client.chat.completions.create.return_value = mock.Mock(
            choices=[mock.Mock(message=message)]
        )
        return service

    def sent_prompt(self, service):
        return service.client.chat.completions.create.call_args.kwargs['messages'][1]['content']

    @override_settings(OPENAI_API_KEY='')
    def test_missing_key_raises(self):
        with self.assertRaises(AIServiceNotConfigured):
            OpenAIService()

    def test_summary_is_parsed_and_input_truncated(self):
        service = self.make_service('{"full_summary": " Text ", "key_points": ["a", "", "b"]}')
        result = service.generate_summary('x' * 500)

        self.assertEqual(result, {'full_summary': 'Text', 'key_points': ['a', 'b']})
        self.assertIn('x' * 50, self.sent_prompt(service))
        self.assertNotIn('x' * 51, self.sent_prompt(service))
        kwargs = service.client.chat.completions.create.call_args.kwargs
        self.assertEqual(kwargs['response_format'], {'type': 'json_object'})

    def test_invalid_responses_raise(self):
        for content in ('not json', '[]', '{"key_points": []}', '{"full_summary": ""}'):
            with self.assertRaises(AIServiceError):
                self.make_service(content).generate_summary('text')

    def test_flashcards_are_cleaned_and_limited(self):
        service = self.make_service(
            '{"flashcards": [{"question": "Q1", "answer": "A1"}, {"question": "", "answer": "x"},'
            ' "junk", {"question": "Q2", "answer": "A2", "category": "Math"},'
            ' {"question": "Q3", "answer": "A3"}]}'
        )
        cards = service.generate_flashcards('text', num_cards=2)
        self.assertEqual(cards, [
            {'question': 'Q1', 'answer': 'A1', 'category': 'General'},
            {'question': 'Q2', 'answer': 'A2', 'category': 'Math'},
        ])

    def test_no_usable_flashcards_raises(self):
        with self.assertRaises(AIServiceError):
            self.make_service('{"flashcards": []}').generate_flashcards('text')
